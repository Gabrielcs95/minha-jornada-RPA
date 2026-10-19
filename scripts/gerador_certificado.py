from pathlib import Path

import docx
import openpyxl

ROOT_DIR = Path(__file__).resolve().parents[1]
certificado = ROOT_DIR / "data" / "input" / "Certificado1.docx"
a_novos = ROOT_DIR / "data" / "input" / "alunosnovos.xlsx"
output_dir = ROOT_DIR / "data" / "output"
output_dir.mkdir(parents=True, exist_ok=True)

#usamos o openpyxl para abrir o arquivo excel
arquivo_excel = openpyxl.load_workbook(a_novos)

#usamos o método active para selecionar a planilha ativa do arquivo excel
arquivo_novo = arquivo_excel.active

#este for vai fazer a leitura de cada linha do arquivo excel, começando da segunda linha (min_row=2) e pegando apenas os valores (values_only=True)
for linha in arquivo_novo.iter_rows(min_row=2, values_only=True):
    nome_atual = linha[0]
    documento = docx.Document(certificado)
    
    #este outro for vai fazer a leitura de cada parágrafo do documento do word, e se encontrar o texto "@nome" ele vai substituir pelo nome atual da linha do excel
    for paragrafo in documento.paragraphs:
        if "@nome" in paragrafo.text:
            paragrafo.text = paragrafo.text.replace("@nome", nome_atual)
        
    documento.save(output_dir / f"Certificado_{nome_atual}.docx")
